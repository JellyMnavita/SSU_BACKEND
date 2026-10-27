"""Word — guide trajets sécurisés et secteurs SSU."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from prediction_shared.guide_content import build_guide_document


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for j, text in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = text
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    for i, row in enumerate(rows, start=1):
        for j, text in enumerate(row):
            table.rows[i].cells[j].text = text


def generate_guide_docx(output: Path, doc: dict[str, Any] | None = None) -> Path:
    doc_data = doc or build_guide_document()
    document = Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)

    title = document.add_heading(doc_data["title"], 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = document.add_paragraph(doc_data["subtitle"])
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = document.add_paragraph(
        f"Généré le {doc_data['generated_at']} — {doc_data['version']}"
    )
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("Chaîne de traitement", level=1)
    for step in doc_data["architecture"]:
        document.add_paragraph(step, style="List Bullet")

    document.add_page_break()

    for section in doc_data["sections"]:
        document.add_heading(f"{section['id']}. {section['title']}", level=1)
        # Retirer balises HTML simples pour Word
        text = (
            section["body"]
            .replace("<br/>", "\n")
            .replace("<br>", "\n")
            .replace("<b>", "")
            .replace("</b>", "")
            .replace("<i>", "")
            .replace("</i>", "")
        )
        document.add_paragraph(text)

    document.add_page_break()
    document.add_heading("Annexe A — Formules clés", level=1)
    for item in doc_data["formulas"]:
        p = document.add_paragraph()
        p.add_run(f"{item['name']} : ").bold = True
        p.add_run(item["expr"])

    for table_def in doc_data["tables"]:
        document.add_heading(table_def["title"], level=2)
        _add_table(document, table_def["headers"], table_def["rows"])

    document.add_page_break()
    document.add_heading("Annexe B — Limites et bonnes pratiques", level=1)
    for limit in doc_data["limits"]:
        document.add_paragraph(limit, style="List Bullet")

    document.save(str(output))
    return output
