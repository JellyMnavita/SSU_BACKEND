"""Génération du rapport Word (.docx) — pipeline NLP SSU."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from nlp_shared.lexicon_data import INCIDENT_TYPES
from nlp_shared.report_assets import NlpReportAssets, prepare_report_assets
from nlp_shared.training.dataset import resolve_nlp_data_dir


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            table.rows[i].cells[j].text = str(cell_text)
            if i == 0:
                for p in table.rows[i].cells[j].paragraphs:
                    for run in p.runs:
                        run.bold = True


def _add_image(doc: Document, path: Path, width_cm: float = 14) -> None:
    if path.is_file():
        doc.add_picture(str(path), width=Cm(width_cm))
        doc.add_paragraph()


def generate_comprehensive_docx(
    output: Path,
    data_dir: Path | None = None,
    assets: NlpReportAssets | None = None,
    *,
    api_url: str | None = None,
) -> Path:
    root = data_dir or resolve_nlp_data_dir()
    assets = assets or prepare_report_assets(root, api_url=api_url)
    profile = assets.profile
    training = assets.training
    metrics = assets.metrics
    comparison = training.get("comparison") or {}
    totals = profile.get("totals") or {}

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    title = doc.add_heading("Rapport NLP — SSU Bokengi", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _add_table(
        doc,
        [
            ["Indicateur", "Valeur"],
            ["Textes corpus", str(totals.get("all", "—"))],
            [
                "Train / Val / Test / Gold",
                f"{totals.get('train')} / {totals.get('val')} / {totals.get('test')} / {totals.get('gold')}",
            ],
            ["Modèle ML retenu", str(training.get("primary_ml") or "—")],
            ["F1 macro (gold)", str(metrics.get("f1_macro", "—"))],
        ],
    )
    doc.add_page_break()

    doc.add_heading("1. Objectif", level=1)
    doc.add_paragraph(
        f"Classification des signalements en {', '.join(INCIDENT_TYPES)}. "
        "Pipeline : filtres anodins → règles lexique → ML (TF-IDF/LSTM/DistilCamemBERT) → fusion vision."
    )

    doc.add_heading("2. Corpus", level=1)
    sources = profile.get("sources") or {}
    doc.add_paragraph(
        f"Total {totals.get('all')} textes. Sources synthetic/humvi/ssu/seed : "
        f"{sources.get('synthetic')}/{sources.get('humvi')}/{sources.get('ssu')}/{sources.get('seed')}."
    )
    _add_image(doc, assets.eda_plots.get("label_distribution", Path()))

    doc.add_heading("3. Sélection du modèle", level=1)
    doc.add_paragraph(training.get("selection_reason") or "Meilleur F1 macro validation.")
    comp_rows = [["Modèle", "Val F1", "Test F1"]]
    for name, entry in comparison.items():
        comp_rows.append([
            name,
            f"{(entry.get('val_f1_macro') or 0):.4f}",
            f"{(entry.get('test_f1_macro') or 0):.4f}",
        ])
    _add_table(doc, comp_rows)
    _add_image(doc, assets.eda_plots.get("model_comparison", Path()))

    doc.add_heading("4. Règles et filtres", level=1)
    doc.add_paragraph(
        f"Règles eval : {assets.rules_eval.get('ok')}/{assets.rules_eval.get('total')}. "
        f"Anodin eval : {assets.benign_eval.get('ok')}/{assets.benign_eval.get('total')}."
    )

    doc.add_heading("5. Gold eval", level=1)
    doc.add_paragraph(
        f"F1 macro {metrics.get('f1_macro')}, accuracy {metrics.get('accuracy')}, "
        f"samples {assets.gold.get('samples')}."
    )
    _add_image(doc, assets.cm_path)
    _add_image(doc, assets.f1_path)

    doc.add_heading("6. Recommandations", level=1)
    doc.add_paragraph(
        "Enrichir le corpus avec signalements validés, réentraîner via l'admin ou "
        "bash scripts/run_nlp_retrain.sh."
    )

    doc.save(str(output))
    return output
