"""Génération du rapport Word (.docx) — pipeline prédiction spatiale SSU."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from prediction_shared.labels import FEATURE_NAMES
from prediction_shared.report_assets import PredictionReportAssets, prepare_report_assets
from prediction_shared.dataset import resolve_prediction_data_dir


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            table.rows[i].cells[j].text = str(cell_text)
            if i == 0:
                for p in table.rows[i].cells[j].paragraphs:
                    for run in p.runs:
                        run.bold = True


def _add_image(doc: Document, path: Path, width_cm: float = 14) -> None:
    if path.is_file():
        doc.add_picture(str(path), width=Cm(width_cm))
        doc.add_paragraph()


def generate_comprehensive_docx(
    output: Path,
    data_dir: Path | None = None,
    assets: PredictionReportAssets | None = None,
) -> Path:
    root = data_dir or resolve_prediction_data_dir()
    assets = assets or prepare_report_assets(root)
    profile = assets.profile
    training = assets.training
    metrics = assets.metrics
    test_metrics = assets.test_metrics
    comparison = training.get("comparison") or {}
    splits = profile.get("splits") or {}
    cleaning = profile.get("cleaning") or {}

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    title = doc.add_heading("Rapport Prédiction spatiale — SSU Bokengi", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _add_table(
        doc,
        [
            ["Indicateur", "Valeur"],
            ["Signalements nettoyés", str(profile.get("total_clean", "—"))],
            [
                "Train / Val / Test / Gold",
                f"{splits.get('train')} / {splits.get('val')} / "
                f"{splits.get('test')} / {splits.get('gold')}",
            ],
            ["Modèle retenu", str(training.get("algorithm") or "—")],
            ["F1 macro (test)", str(test_metrics.get("f1_macro", "—"))],
            ["F1 macro (gold)", str(metrics.get("f1_macro", "—"))],
        ],
    )
    doc.add_page_break()

    doc.add_heading("1. Objectif", level=1)
    doc.add_paragraph(
        "Prédire le degré de danger (faible, modéré, élevé) à partir de la géolocalisation "
        "et du contexte spatio-temporel. Alimente la carte de chaleur et prediction-service."
    )

    doc.add_heading("2. Nettoyage des données", level=1)
    doc.add_paragraph(
        f"Lignes lues {cleaning.get('rows_read')}, conservées {cleaning.get('rows_kept')}. "
        f"Rejets : coords invalides={cleaning.get('removed_invalid_coords')}, "
        f"hors Kinshasa={cleaning.get('removed_out_of_bbox')}, "
        f"degré invalide={cleaning.get('removed_invalid_degree')}, "
        f"doublons={cleaning.get('removed_duplicates')}."
    )
    _add_image(doc, assets.eda_plots.get("class_balance", Path()))

    doc.add_heading("3. Analyse exploratoire", level=1)
    for key in ("spatial_scatter", "hour_distribution", "weekday_distribution", "top_locations"):
        _add_image(doc, assets.eda_plots.get(key, Path()))

    doc.add_heading("4. Features", level=1)
    doc.add_paragraph("Variables : " + ", ".join(FEATURE_NAMES) + ".")
    doc.add_paragraph(
        "Densité 500 m et degré moyen voisins calculés en leave-one-out (anti-fuite)."
    )
    for step in assets.eda_summary.get("preprocessing") or []:
        doc.add_paragraph(step, style="List Bullet")

    doc.add_heading("5. Validation croisée et sélection", level=1)
    doc.add_paragraph(training.get("selection_reason") or "Meilleur F1 macro en CV 5-fold.")
    comp_rows = [["Modèle", "F1 macro", "Std", "Accuracy"]]
    for name, entry in comparison.items():
        comp_rows.append([
            name,
            f"{entry.get('f1_macro_mean', 0):.4f}",
            f"±{entry.get('f1_macro_std', 0):.4f}",
            f"{entry.get('accuracy_mean', 0):.4f}",
        ])
    _add_table(doc, comp_rows)
    _add_image(doc, assets.cv_path)

    doc.add_heading("6. Résultats test et gold_eval", level=1)
    doc.add_paragraph(
        f"Test — accuracy={test_metrics.get('accuracy')}, "
        f"f1_macro={test_metrics.get('f1_macro')}. "
        f"Gold — accuracy={metrics.get('accuracy')}, f1_macro={metrics.get('f1_macro')}."
    )
    _add_image(doc, assets.cm_path)
    _add_image(doc, assets.f1_path)

    doc.add_heading("7. Recommandations", level=1)
    doc.add_paragraph(
        "Enrichir les signalements validés, surveiller le déséquilibre des classes, "
        "réentraîner via bash scripts/run_prediction_retrain.sh."
    )

    doc.save(str(output))
    return output
